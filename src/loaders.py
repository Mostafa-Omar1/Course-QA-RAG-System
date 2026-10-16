"""Extract and clean text from PDF, DOCX, CSV and TXT course files.

Every loader returns a list of "documents": dicts with `text` and `metadata`
(course, source file, file type and a location such as page or row range).
"""
from pathlib import Path
import re

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".csv", ".txt"}


def clean_text(text: str) -> str:
    """Normalize whitespace and remove common extraction artifacts."""
    text = text.replace("\x00", " ")
    text = re.sub(r"-\n(\w)", r"\1", text)          # re-join hyphenated line breaks
    text = re.sub(r"[ \t ]+", " ", text)        # collapse spaces/tabs
    text = re.sub(r" *\n *", "\n", text)             # trim spaces around newlines
    text = re.sub(r"\n{3,}", "\n\n", text)           # max one blank line
    return text.strip()


def _meta(path: Path, course: str, **extra) -> dict:
    return {"course": course, "source": path.name, "file_type": path.suffix.lstrip(".").lower(), **extra}


def load_pdf(path: Path, course: str) -> list[dict]:
    docs = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            docs.append({"text": text, "metadata": _meta(path, course, location=f"page {i}")})
    return docs


def load_docx(path: Path, course: str) -> list[dict]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = clean_text("\n".join(parts))
    return [{"text": text, "metadata": _meta(path, course, location="document")}] if text else []


def load_csv(path: Path, course: str, rows_per_doc: int = 10) -> list[dict]:
    """Turn each row into a 'column: value' sentence so it embeds well, grouped in small batches."""
    df = pd.read_csv(path).fillna("")
    docs = []
    for start in range(0, len(df), rows_per_doc):
        block = df.iloc[start:start + rows_per_doc]
        lines = ["; ".join(f"{col}: {row[col]}" for col in df.columns) for _, row in block.iterrows()]
        text = clean_text("\n".join(lines))
        loc = f"rows {start + 1}-{start + len(block)}"
        docs.append({"text": text, "metadata": _meta(path, course, location=loc)})
    return docs


def load_txt(path: Path, course: str) -> list[dict]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    return [{"text": text, "metadata": _meta(path, course, location="document")}] if text else []


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".csv": load_csv, ".txt": load_txt}


def load_file(path: Path, course: str) -> list[dict]:
    return LOADERS[path.suffix.lower()](path, course)


def load_courses(root: Path) -> list[dict]:
    """Each sub-folder of `root` is one course; load every supported file inside it."""
    docs = []
    for course_dir in sorted(p for p in root.iterdir() if p.is_dir()):
        for path in sorted(course_dir.rglob("*")):
            if path.suffix.lower() in SUPPORTED:
                docs.extend(load_file(path, course_dir.name))
    return docs
