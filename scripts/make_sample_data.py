"""Generate sample course materials for 3 courses in PDF, DOCX, CSV and TXT formats.

Run once:  python scripts/make_sample_data.py
Output:    data/courses/<course>/...
"""
from pathlib import Path
import csv

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent / "data" / "courses"


# ---------------------------------------------------------------------------
# Course 1: Machine Learning Fundamentals
# ---------------------------------------------------------------------------
ML_PDF = [
    ("Lecture 1: What is Machine Learning?", [
        "Machine learning (ML) is a field of artificial intelligence in which computers learn patterns "
        "from data instead of being explicitly programmed with rules. A model is trained on examples and "
        "then used to make predictions on new, unseen data.",
        "There are three main types of machine learning. Supervised learning uses labeled data, where each "
        "example has an input and a known output. Unsupervised learning finds structure in unlabeled data, "
        "for example grouping customers into clusters. Reinforcement learning trains an agent to take "
        "actions in an environment so as to maximize a cumulative reward.",
        "Typical supervised tasks are classification, where the output is a category such as spam or not "
        "spam, and regression, where the output is a continuous number such as a house price.",
    ]),
    ("Lecture 2: Linear Regression", [
        "Linear regression models the relationship between input features and a continuous target as a "
        "straight line: y = w * x + b, where w is the weight (slope) and b is the bias (intercept).",
        "The model is trained by minimizing the Mean Squared Error (MSE), which is the average of the "
        "squared differences between predicted and actual values.",
        "Gradient descent is an optimization algorithm that updates the parameters in the direction of the "
        "negative gradient of the loss. The learning rate controls the size of each step: a learning rate "
        "that is too large can diverge, while one that is too small converges very slowly.",
    ]),
    ("Lecture 3: Overfitting and Evaluation", [
        "Overfitting happens when a model learns the training data too well, including its noise, and "
        "performs poorly on new data. Underfitting happens when a model is too simple to capture the "
        "underlying pattern.",
        "To detect overfitting we split the data into a training set, a validation set and a test set. "
        "A common split is 70% training, 15% validation and 15% test. K-fold cross-validation splits the "
        "data into k folds and trains k times, each time holding out a different fold.",
        "Regularization reduces overfitting by penalizing large weights. L1 regularization (Lasso) can "
        "drive some weights to exactly zero, while L2 regularization (Ridge) shrinks all weights smoothly.",
        "For classification, accuracy is the fraction of correct predictions. Precision is TP / (TP + FP), "
        "recall is TP / (TP + FN), and the F1 score is the harmonic mean of precision and recall.",
    ]),
]

ML_TXT = """Machine Learning Fundamentals - Course Syllabus

Instructor: Dr. Sara Ahmed
Duration: 8 weeks
Grading: Assignments 40%, Midterm exam 25%, Final project 35%.

Week 1: Introduction to machine learning and types of learning.
Week 2: Linear regression and gradient descent.
Week 3: Logistic regression and classification metrics.
Week 4: Overfitting, regularization and cross-validation.
Week 5: Decision trees and random forests.
Week 6: K-means clustering and PCA.
Week 7: Introduction to neural networks.
Week 8: Final project presentations.

Late submission policy: assignments submitted late lose 10% per day, up to a maximum of 3 days.
After 3 days the assignment receives zero.

Office hours: Sundays and Tuesdays from 2 PM to 4 PM.
"""

ML_CSV = [
    ("term", "definition"),
    ("Feature", "An input variable used by the model to make a prediction."),
    ("Label", "The known output or target value in supervised learning."),
    ("Epoch", "One complete pass of the training algorithm over the whole training dataset."),
    ("Batch size", "The number of training examples processed before the model parameters are updated."),
    ("Random forest", "An ensemble of many decision trees whose predictions are averaged or voted on."),
    ("K-means", "A clustering algorithm that partitions data into k clusters by minimizing distance to cluster centroids."),
    ("PCA", "Principal Component Analysis, a dimensionality reduction technique that projects data onto directions of maximum variance."),
    ("Activation function", "A non-linear function such as ReLU or sigmoid applied to a neuron's output."),
]

# ---------------------------------------------------------------------------
# Course 2: Python Programming
# ---------------------------------------------------------------------------
PY_DOCX = [
    ("Chapter 1: Python Basics", [
        "Python is a high-level, interpreted programming language known for its readable syntax. "
        "Indentation is used to define code blocks instead of curly braces.",
        "Python's basic data types are int, float, str and bool. Variables do not need type declarations; "
        "the type is determined at runtime (dynamic typing).",
        "The built-in function type() returns the type of an object, and len() returns the length of a sequence.",
    ]),
    ("Chapter 2: Data Structures", [
        "A list is an ordered, mutable collection written with square brackets, for example [1, 2, 3]. "
        "A tuple is ordered but immutable and is written with parentheses, for example (1, 2, 3).",
        "A dictionary stores key-value pairs, for example {'name': 'Ali', 'age': 21}. Keys must be immutable "
        "and unique. A set is an unordered collection of unique elements.",
        "List comprehensions provide a concise way to create lists: [x * x for x in range(10) if x % 2 == 0] "
        "creates the squares of even numbers below 10.",
    ]),
    ("Chapter 3: Functions and Error Handling", [
        "Functions are defined with the def keyword. *args collects extra positional arguments into a tuple, "
        "and **kwargs collects extra keyword arguments into a dictionary.",
        "A lambda is a small anonymous function, for example lambda x: x + 1.",
        "Exceptions are handled with try / except blocks. The finally block always runs, whether an "
        "exception occurred or not, and is commonly used to release resources such as files.",
        "The with statement uses a context manager to automatically open and close resources, for example "
        "with open('data.txt') as f: content = f.read().",
    ]),
]

PY_TXT = """Python Programming - Course Information

Instructor: Eng. Omar Hassan
Duration: 6 weeks
Grading: Weekly labs 50%, Final exam 50%.

Required software: Python 3.10 or newer and Visual Studio Code.

Lab rules: Each weekly lab must be submitted on the course portal before Thursday at 11:59 PM.
Code must follow PEP 8, the official Python style guide (4 spaces for indentation, snake_case for
function and variable names, lines no longer than 79 characters).

Virtual environments: create one with "python -m venv .venv" and install packages with pip.
"""

PY_CSV = [
    ("week", "topic", "lab_assignment"),
    ("1", "Variables, data types and input/output", "Build a simple calculator"),
    ("2", "Conditionals and loops", "FizzBuzz and number guessing game"),
    ("3", "Lists, tuples, dictionaries and sets", "Student grade book"),
    ("4", "Functions, *args and **kwargs", "Text statistics utility"),
    ("5", "File handling and exceptions", "CSV contact manager"),
    ("6", "Object-oriented programming: classes and inheritance", "Bank account system"),
]

# ---------------------------------------------------------------------------
# Course 3: Database Systems
# ---------------------------------------------------------------------------
DB_PDF = [
    ("Lecture 1: Relational Databases", [
        "A relational database stores data in tables made of rows (records) and columns (attributes). "
        "A primary key uniquely identifies each row of a table, and a foreign key is a column that "
        "references the primary key of another table.",
        "SQL (Structured Query Language) is the standard language for relational databases. SELECT reads "
        "data, INSERT adds rows, UPDATE modifies rows, and DELETE removes rows.",
    ]),
    ("Lecture 2: Normalization", [
        "Normalization organizes tables to reduce redundancy and avoid update anomalies.",
        "First Normal Form (1NF) requires that every column holds atomic values and there are no repeating groups. "
        "Second Normal Form (2NF) requires 1NF and that every non-key attribute depends on the whole primary key, "
        "not just part of it. Third Normal Form (3NF) requires 2NF and that there are no transitive dependencies, "
        "meaning non-key attributes depend only on the key.",
    ]),
    ("Lecture 3: Transactions and ACID", [
        "A transaction is a sequence of operations executed as a single logical unit of work.",
        "ACID stands for Atomicity, Consistency, Isolation and Durability. Atomicity means all operations of a "
        "transaction succeed or none do. Consistency means a transaction moves the database from one valid state "
        "to another. Isolation means concurrent transactions do not interfere with each other. Durability means "
        "once a transaction is committed it survives system failures.",
        "An index is a data structure, usually a B-tree, that speeds up data retrieval at the cost of extra "
        "storage and slower writes.",
    ]),
]

DB_DOCX = [
    ("SQL Joins Guide", [
        "An INNER JOIN returns only the rows that have matching values in both tables.",
        "A LEFT JOIN returns all rows from the left table and the matching rows from the right table; "
        "unmatched right-side columns are filled with NULL. A RIGHT JOIN is the mirror image.",
        "A FULL OUTER JOIN returns all rows when there is a match in either table.",
        "GROUP BY groups rows that share values so aggregate functions such as COUNT, SUM and AVG can be applied. "
        "HAVING filters groups after aggregation, whereas WHERE filters rows before aggregation.",
    ]),
]

DB_CSV = [
    ("assessment", "weight", "due_week", "notes"),
    ("Quiz 1", "10%", "3", "Covers relational model and basic SQL"),
    ("Midterm exam", "30%", "6", "Closed book; covers SQL, joins and normalization"),
    ("Database design project", "25%", "10", "Design an ER diagram and normalized schema for a library system"),
    ("Final exam", "35%", "12", "Comprehensive; includes transactions, ACID and indexing"),
]

DB_TXT = """Database Systems - Course Overview

Instructor: Dr. Mona Khaled
Duration: 12 weeks
Database engine used in labs: PostgreSQL 16.

Attendance policy: students must attend at least 75% of lectures to be allowed to sit the final exam.

Recommended textbook: "Database System Concepts" by Silberschatz, Korth and Sudarshan.
"""


def write_pdf(path: Path, title: str, sections):
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for i, (heading, paras) in enumerate(sections):
        if i:
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story.append(Paragraph(p, styles["BodyText"]))
            story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def write_docx(path: Path, title: str, sections):
    doc = Document()
    doc.add_heading(title, 0)
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p)
    doc.save(path)


def write_csv(path: Path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def main():
    courses = {
        "Machine Learning Fundamentals": [
            ("ml_lectures.pdf", write_pdf, ML_PDF),
            ("ml_syllabus.txt", None, ML_TXT),
            ("ml_glossary.csv", write_csv, ML_CSV),
        ],
        "Python Programming": [
            ("python_notes.docx", write_docx, PY_DOCX),
            ("python_course_info.txt", None, PY_TXT),
            ("python_schedule.csv", write_csv, PY_CSV),
        ],
        "Database Systems": [
            ("db_lectures.pdf", write_pdf, DB_PDF),
            ("sql_joins_guide.docx", write_docx, DB_DOCX),
            ("db_assessments.csv", write_csv, DB_CSV),
            ("db_overview.txt", None, DB_TXT),
        ],
    }
    for course, files in courses.items():
        folder = ROOT / course
        folder.mkdir(parents=True, exist_ok=True)
        for name, writer, content in files:
            path = folder / name
            if writer is write_csv:
                writer(path, content)
            elif writer is None:
                path.write_text(content, encoding="utf-8")
            else:
                writer(path, course, content)
            print(f"wrote {path.relative_to(ROOT.parent.parent)}")


if __name__ == "__main__":
    main()
